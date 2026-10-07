# Genera un PDF y un DOCX con texto normal y un cuadro (tabla) para probar la lectura ordenada de cuadros.
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import docx
pdfmetrics.registerFont(TTFont('Serif', '/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf'))
pdfmetrics.registerFont(TTFont('Sans', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
st = getSampleStyleSheet(); st['Normal'].fontName = 'Serif'; st['Normal'].fontSize = 11; st['Normal'].leading = 15
ANTES = 'El consumo de energía eléctrica cambió de forma desigual entre sectores durante el último año, según el informe de la empresa distribuidora.'
DESPUES = 'Como se observa, el sector industrial sigue siendo el mayor consumidor, aunque el residencial fue el que más creció.'
FILAS = [['Sector', 'Consumo (GWh)', 'Variación', 'Usuarios'],
         ['Residencial', '1.250', '6,5 %', '2.300.000'],
         ['Comercial', '840', '3,1 %', '310.000'],
         ['Industrial', '2.100', '1,2 %', '12.500'],
         ['Oficial', '320', '-0,8 %', '4.100']]
doc = SimpleDocTemplate('tablas.pdf', pagesize=letter)
t = Table(FILAS, colWidths=[110, 110, 90, 100])
t.setStyle(TableStyle([('FONT', (0, 0), (-1, -1), 'Sans', 10), ('GRID', (0, 0), (-1, -1), 0.5, colors.grey), ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey)]))
doc.build([Paragraph(ANTES, st['Normal']), Spacer(1, 14), Paragraph('Tabla 1. Consumo de energía por sector, 2025', st['Normal']), Spacer(1, 6), t, Spacer(1, 14), Paragraph(DESPUES, st['Normal'])])
d = docx.Document()
d.add_paragraph(ANTES)
d.add_paragraph('Tabla 1. Consumo de energía por sector, 2025')
tb = d.add_table(rows=len(FILAS), cols=4)
for i, f in enumerate(FILAS):
    for j, c in enumerate(f): tb.cell(i, j).text = c
d.add_paragraph(DESPUES)
d.save('tablas.docx')
print('ok')
